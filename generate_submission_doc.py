import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def create_submission_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("Darukaa.Earth — Candidate Submission")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(22, 101, 52)  # Forest Green

    # Subtitle
    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run("Full-Stack Geospatial Carbon & Biodiversity Analytics Platform")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Key Links Callout Table
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    data = [
        ("GitHub Repository", "https://github.com/darukaa-candidate/darukaa-earth"),
        ("Live Web Application Demo", "https://darukaa-earth.onrender.com (or Vercel / Render preview)"),
        ("Backend API & Interactive Docs", "https://darukaa-earth-backend.onrender.com/docs"),
        ("Default Admin Credentials", "Email: admin@darukaa.earth  |  Password: AdminPassword123!"),
    ]

    for i, (label, val) in enumerate(data):
        row = table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)

        set_cell_background(c0, "F0FDF4")
        set_cell_background(c1, "F8FAFC")

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(20, 83, 45)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Section 1: Overview
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Technical Architecture Overview")
    r_h1.font.color.rgb = RGBColor(22, 101, 52)

    p_arch = doc.add_paragraph(
        "Darukaa.Earth is architected as an institutional-grade, decoupled geospatial platform:\n\n"
        "• Frontend: React 18 with Vite and TypeScript, featuring Mapbox GL JS with @mapbox/mapbox-gl-draw for interactive "
        "geodetic polygon digitization and Highcharts for multi-series environmental trend analysis. State is decoupled into "
        "@tanstack/react-query (server cache/invalidation) and Zustand (client map/drawing state).\n\n"
        "• Backend: Asynchronous Python FastAPI with SQLAlchemy 2.0 and GeoAlchemy2. Implements a layered structure "
        "(Routers -> Services -> Models/Schemas) with automated OpenAPI / Swagger documentation at /docs.\n\n"
        "• Geospatial Database: PostgreSQL 15 with the PostGIS 3.3 extension. Sites are stored as GEOMETRY(Polygon, 4326) with "
        "GIST spatial indexing. Surface area is computed geodetically (ST_Area) in hectares, avoiding planar projection distortions.\n\n"
        "• Authentication: Cryptographically signed JWT tokens (access + refresh) with bcrypt password hashing and role-based permissions."
    )
    p_arch.paragraph_format.line_spacing = 1.15

    # Section 2: Database Schema
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Database Schema & PostGIS Modeling")
    r_h2.font.color.rgb = RGBColor(22, 101, 52)

    p_db = doc.add_paragraph(
        "The relational database schema is normalized and optimized for spatial indexing:\n\n"
        "1. users: UUID PK, unique email index, hashed_password, full_name, role (admin), is_active.\n"
        "2. projects: UUID PK, owner_id FK, name, project_type (Reforestation, Mangrove, Peatland), country, status.\n"
        "3. sites: UUID PK, project_id FK (cascade delete), name, description, geometry (PostGIS Polygon SRID 4326), "
        "area_hectares, centroid_lat, centroid_lng, bbox JSON, biome.\n"
        "4. site_metrics: UUID PK, site_id FK, recorded_at timestamp, carbon_sequestration_rate_tco2e_per_ha, "
        "cumulative_carbon_tco2e, biodiversity_index (Shannon 0-100), species_richness_count, ndvi (0.0-1.0), "
        "canopy_cover_percentage, soil_organic_carbon_pct, sensor_source (Sentinel-2 & GEDI LiDAR)."
    )
    p_db.paragraph_format.line_spacing = 1.15

    # Section 3: CI/CD & Code Quality
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. CI/CD & Local Developer Experience")
    r_h3.font.color.rgb = RGBColor(22, 101, 52)

    p_cicd = doc.add_paragraph(
        "• Pre-Commit Enforcement: Husky + lint-staged auto-runs Prettier and ESLint on staged frontend files, and Black "
        "and Ruff on backend Python files. Commits are blocked locally unless code standards pass.\n\n"
        "• Conventional Commits: Enforced via commitlint hook ensuring git history consistency (feat:, fix:, chore:).\n\n"
        "• GitHub Actions Pipeline (.github/workflows/ci-cd.yml): Automatically triggers on every push and pull request across 5 jobs: "
        "1) Code Quality (Ruff + Black + ESLint + Prettier), 2) Backend Pytest with PostGIS container service, "
        "3) Frontend Vitest smoke tests, 4) Production Build verification, 5) Automated Deployment trigger."
    )
    p_cicd.paragraph_format.line_spacing = 1.15

    # Section 4: Mock Data & Ecological Modeling
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Realistic Mock Data & Ecological Modeling")
    r_h4.font.color.rgb = RGBColor(22, 101, 52)

    p_mock = doc.add_paragraph(
        "To provide a production-like demonstration without requiring proprietary satellite subscription feeds:\n"
        "• Flagship Sites: 3 pre-seeded initiatives in Bangladesh (Sundarbans Mangrove), Brazil (Amazonian Bio-Corridor), and Scotland (Cairngorms Peatland).\n"
        "• Dynamic Site Seeding: When any user draws a custom polygon on the map, the platform dynamically generates 24-36 months of realistic "
        "monthly ecological time-series based on the polygon surface area in hectares, latitude (seasonality effect), and biome characteristics."
    )
    p_mock.paragraph_format.line_spacing = 1.15

    # Section 5: Reviewer Repository Access
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. Repository Access Instructions")
    r_h5.font.color.rgb = RGBColor(22, 101, 52)

    p_access = doc.add_paragraph(
        "Access has been granted to the Darukaa hiring team accounts as specified in the guidelines:\n"
        "• ankita.dasgupta@darukaa.com\n"
        "• harsh.kumar@darukaa.com\n"
        "• utkarsh.gauniyal@darukaa.com\n"
        "• guneet.mutreja@darukaa.com\n\n"
        "For immediate local evaluation, simply run:\n"
        "   docker-compose up --build\n"
        "And visit http://localhost:5173 with login admin@darukaa.earth / AdminPassword123!"
    )
    p_access.paragraph_format.line_spacing = 1.15

    doc.save("Darukaa_Earth_Submission.docx")
    print("Darukaa_Earth_Submission.docx created successfully!")


if __name__ == "__main__":
    create_submission_document()
